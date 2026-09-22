#!/usr/bin/env python3
"""Write the Spate starter deck, `website/static/brand/spate-slides.pptx`.

Title, section, content, code and closing slides, each on the light and the
dark ground. Layout is in pixels on a 1600×900 canvas, one pixel being 1/120
in on the 13.333×7.5 in slide. The wordmark and the carriers are PNG pictures
rendered by resvg at twice their size; everything else is an editable text box
set in IBM Plex Sans or Overpass Mono, which the presenter installs. Reruns
write the same bytes.

Needs resvg on PATH.
"""
import datetime
import io
import pathlib
import subprocess
import tempfile
import zipfile

from fontTools.misc.transform import Transform
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

import brandgen
import waterline
from brandgen import TOKENS
from waterline import MONO_FONT, SRC_FONT

OUT = brandgen.STATIC.parent / "brand" / "spate-slides.pptx"
W, H = 1600, 900
PX = 7620  # EMU per canvas pixel: 914400 EMU per inch, 120 px per inch
FIXED = datetime.datetime(2026, 1, 1)

# Office finds a static cut by its legacy family name. IBM's own release names
# the Medium cut "IBM Plex Sans Medm"; "IBM Plex Sans Medium" falls back to
# another face.
SANS, SANS_MEDIUM = "IBM Plex Sans", "IBM Plex Sans Medm"
MONO, MONO_MEDIUM = "Overpass Mono", "Overpass Mono Medium"
# Baseline to the top of the line box, per em: the fonts' hhea ascent.
ASCENT = {
    face: TTFont(path)["hhea"].ascent / TTFont(path)["head"].unitsPerEm
    for face, path in ((SANS, SRC_FONT), (SANS_MEDIUM, SRC_FONT), (MONO, MONO_FONT), (MONO_MEDIUM, MONO_FONT))
}


def rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip("#"))


def emu(px):
    return Emu(round(px * PX))


class Deck:
    def __init__(self, art, scratch):
        self.art = art
        self.scratch = scratch
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = emu(W), emu(H)
        self.blank = self.prs.slide_layouts[6]

    def slide(self, c):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(c["bg"])
        return s

    def picture(self, s, name, svg, x, y, w):
        """Render `svg` at twice `w` and place it with its top-left at (x, y), `w` wide."""
        src, png = self.scratch / f"{name}.svg", self.scratch / f"{name}.png"
        src.write_text(svg)
        subprocess.run(["resvg", "--width", str(2 * w), str(src), str(png)], check=True)
        s.shapes.add_picture(str(png), emu(x), emu(y), width=emu(w))

    def wordmark(self, s, c, ground, x, y, w):
        fx0, fy0, fx1, fy1 = brandgen.finished_ink(self.art)
        k = w / (fx1 - fx0)
        h = (fy1 - fy0) * k
        t = Transform(k, 0, 0, k, -fx0 * k, -fy0 * k)
        body = brandgen.word_paths(self.art, t, *brandgen.two_tone(ground))
        self.picture(s, f"wordmark-{ground}-{w}", brandgen.svg(f"0 0 {w} {h:.2f}", body, (w, h)), x, y, w)

    def carrier(self, s, c, ground, value, box):
        x, y, w, h = box
        above, below = waterline.carrier(value, (0, 0, w, h))
        body = brandgen.pieces([above], [below], Transform(), *brandgen.two_tone(ground))
        self.picture(s, f"carrier-{ground}-{value}-{w}", brandgen.svg(f"0 0 {w} {h}", body, (w, h), value), x, y, w)

    def text(self, s, lines, x, baseline, size, color, face=SANS, tracking=0.0, leading=None,
             width=None, align=PP_ALIGN.LEFT):
        """A text box whose first baseline sits at `baseline`, one paragraph per line."""
        leading = leading or size * 1.2
        top = baseline - ASCENT[face] * size
        width = width or W - x - 80
        box = s.shapes.add_textbox(emu(x), emu(top), emu(width), emu(leading * len(lines)))
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = Pt(leading * 0.6)
            run = p.add_run()
            run.text = line
            run.font.name = face
            run.font.size = Pt(size * 0.6)
            run.font.color.rgb = rgb(color)
            if tracking:
                # Character spacing in hundredths of a point.
                run.font._rPr.set("spc", str(round(tracking * size * 0.6 * 100)))
        return box

    def rule(self, s, c, x, y, w):
        line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x), emu(y), emu(x + w), emu(y))
        line.line.color.rgb = rgb(c["border"])
        line.line.width = emu(1)

    def rect(self, s, x, y, w, h, fill, rounded=False):
        shape = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, emu(x), emu(y), emu(w), emu(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if rounded:
            shape.adjustments[0] = 8 / min(w, h)
        return shape

    def footer(self, s, c, y, left=None):
        self.rule(s, c, 80, y, 1440)
        if left:
            self.text(s, [left], 80, y + 62, 18, c["accent"], MONO_MEDIUM)
        self.text(s, ["spate.kainth.dev"], 1000, y + 62, 17, c["muted"], MONO, width=520, align=PP_ALIGN.RIGHT)

    # ------------------------------------------------------------ slides

    def title(self, ground):
        c = TOKENS[ground]
        s = self.slide(c)
        self.wordmark(s, c, ground, 80, 55, 188)
        self.text(s, ["EVENT OR TRACK"], 80, 228, 19, c["accent"], MONO_MEDIUM)
        self.text(s, ["Title of the talk,", "on two lines."], 74, 353, 97, c["ink"], SANS_MEDIUM, -0.038, 105, width=980)
        self.text(s, ["One sentence the audience leaves with."], 80, 566, 34, c["muted"], width=980)
        self.carrier(s, c, ground, "04", (1070, 205, 450, 397))
        self.rule(s, c, 80, 739, 1440)
        self.text(s, ["Presenter name"], 80, 796, 23, c["ink"], SANS_MEDIUM)
        self.text(s, ["Role, organization"], 80, 834, 21, c["muted"])
        self.text(s, ["spate.kainth.dev"], 1000, 814, 19, c["muted"], MONO, width=520, align=PP_ALIGN.RIGHT)

    def section(self, ground):
        c = TOKENS[ground]
        s = self.slide(c)
        self.wordmark(s, c, ground, 80, 55, 155)
        self.carrier(s, c, ground, "02", (80, 226, 472, 404))
        self.text(s, ["SECTION NAME"], 674, 302, 18, c["accent"], MONO_MEDIUM)
        self.text(s, ["Section heading", "on two lines."], 669, 407, 80, c["ink"], SANS_MEDIUM, -0.025, 91, width=851)
        self.text(s, ["One line on what this section covers."], 674, 584, 29, c["muted"], width=846)
        self.footer(s, c, 761)

    def content(self, ground):
        c = TOKENS[ground]
        s = self.slide(c)
        self.text(s, ["SECTION NAME"], 80, 118, 18, c["accent"], MONO_MEDIUM)
        self.text(s, ["A slide heading on one line"], 76, 206, 64, c["ink"], SANS_MEDIUM, -0.025)
        self.text(
            s,
            ["The first point, in one sentence.", "The second point, in one sentence.",
             "The third point, in one sentence."],
            80, 318, 34, c["ink"], leading=62,
        )
        self.footer(s, c, 761, "02")

    def code(self, ground):
        c = TOKENS[ground]
        s = self.slide(c)
        self.text(s, ["SECTION NAME"], 80, 118, 18, c["accent"], MONO_MEDIUM)
        self.text(s, ["A heading for the code"], 76, 206, 64, c["ink"], SANS_MEDIUM, -0.025)
        self.rect(s, 80, 262, 1440, 440, c["code-bg"], rounded=True)
        self.text(
            s,
            ["// Replace with the code this slide discusses.", "fn main() {",
             '    println!("one pipeline, four stages");', "}"],
            128, 336, 30, c["code-ink"], MONO, leading=48, width=1344,
        )
        self.footer(s, c, 761, "02")

    def closing(self, ground):
        c = TOKENS[ground]
        s = self.slide(c)
        self.wordmark(s, c, ground, 560, 300, 480)
        self.text(s, ["spate.kainth.dev"], 80, 560, 34, c["ink"], MONO, width=1440, align=PP_ALIGN.CENTER)
        self.text(s, ["github.com/spate-etl/spate"], 80, 614, 26, c["muted"], MONO, width=1440,
                  align=PP_ALIGN.CENTER)

    def build(self):
        for ground in ("light", "dark"):
            for make in (self.title, self.section, self.content, self.code, self.closing):
                make(ground)
        props = self.prs.core_properties
        props.title = "Spate slides"
        props.author = props.last_modified_by = "Spate"
        props.created = props.modified = props.last_printed = FIXED
        props.revision = 1
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def stable_zip(blob):
    """Rewrite a zip with fixed entry timestamps, keeping the entry order."""
    src = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            dst.writestr(entry, src.read(info.filename))
    return out.getvalue()


def main():
    brandgen.ensure_fonts()
    art = waterline.artwork(waterline.P)
    with tempfile.TemporaryDirectory() as scratch:
        blob = Deck(art, pathlib.Path(scratch)).build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_bytes(stable_zip(blob))
    print(f"  {OUT.relative_to(brandgen.STATIC.parent.parent)}")


if __name__ == "__main__":
    main()
